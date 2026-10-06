"""ONE-TIME: copy everything already in EventAgent.xlsx into the Google Sheet (same tabs, same columns).
   python migrate_xlsx_to_gsheet.py
Run it once, before starting app.py. A tab that already has data in the Google Sheet is skipped."""
import config, excel_store as g
from openpyxl import load_workbook

wb = load_workbook(config.EXCEL_FILE)
g._ensure()                                              # creates tabs + headers in the Google Sheet
for name in g.SHEETS:
    if name not in wb.sheetnames:
        continue
    if len(g._values(name, fresh=True)) > 1:
        print(f"{name}: Google Sheet already has data - skipped")
        continue
    src = list(wb[name].iter_rows(values_only=True))
    head = [h for h in src[0] if h]
    data = [{h: g._s(r[i]) for i, h in enumerate(head)} for r in src[1:]
            if any(v not in (None, "") for v in r[:len(head)])]
    for i in range(0, len(data), 100):
        g._call("bulk_add", sheet=name, rows=data[i:i + 100])
    print(f"{name}: copied {len(data)} rows")
print("Done.")